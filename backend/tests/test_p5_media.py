"""P5 media tests (offline, hermetic).

- image: structured describe (sections), legacy describe_fn contract,
  failure-proof metadata-only, ingest with sectioned chunks + usage
- audio: deterministic chapterize (gaps/merge/titles, no diarization),
  dispatcher chapter elements, window section stamping, transcript
  chapters, ingest carries sections
- excel relations: detection (shared-column / value-overlap / negatives),
  summary + parse_meta wiring, JOIN validation, deterministic JOIN
  end-to-end over DuckDB
"""

import io
import uuid

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app.common.base import Base
from app.core.security import hash_password
from app.knowledge.models import Chunk, KnowledgeBase, PageSegment, Source
from app.users.models import User
from app.workspaces.models import Workspace, WorkspaceMember
from workers.tasks import run_ingest

engine = create_engine(
    "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
)
TestingSession = sessionmaker(bind=engine, autoflush=False, autocommit=False)
Base.metadata.create_all(engine)


@pytest.fixture()
def db():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    s = TestingSession()
    try:
        yield s
    finally:
        s.close()


def _seed_ws(db, email="u@x.com"):
    pw = hash_password("password123")
    u = User(id=uuid.uuid4(), email=email, password_hash=pw)
    db.add(u)
    w = uuid.uuid4()
    db.add(Workspace(id=w, name="W", type="shared", owner_user_id=u.id))
    db.add(WorkspaceMember(workspace_id=w, user_id=u.id, role="owner"))
    kb = KnowledgeBase(id=uuid.uuid4(), workspace_id=w, title="KB",
                       scope="workspace")
    db.add(kb)
    db.commit()
    return u, w, kb


class FakeStorage:
    def __init__(self):
        self.objects: dict[str, bytes] = {}

    def put(self, key: str, data: bytes, content_type: str = "") -> str:
        self.objects[key] = data
        return key

    def get(self, key: str) -> bytes:
        return self.objects[key]

    def delete(self, key: str) -> None:
        self.objects.pop(key, None)


def _png() -> bytes:
    from PIL import Image

    im = Image.new("RGB", (320, 200), "blue")
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def fake_embed(texts):
    assert texts, "should not embed zero chunks"
    return [[float(len(t) % 7)] * 1536 for t in texts]


# ---------- image ----------

def test_image_legacy_describe_fn_contract():
    from app.knowledge.parsers.fallback import parse_image_meta

    doc = parse_image_meta(_png(), "pic.png",
                           describe_fn=lambda blob: "یک گربه روی مبل نشسته است.")
    assert "320x200" in doc.pages[0].text
    assert "گربه" in doc.pages[0].text
    assert doc.parse_meta["parser"] == "image-meta"
    assert "usage" not in doc.parse_meta
    # sections flow into elements like every other source
    assert any(e.kind == "paragraph" for e in doc.elements)


def test_image_structured_default_path_with_usage(monkeypatch):
    from app.knowledge.parsers import fallback as F

    seen = {}

    def fake_describe(self, blob, prompt, model=None, **kw):
        seen["model"] = model
        assert "صحنه" in prompt  # structured prompt, not free chat
        return ("## صحنه\n\nیک خیابان شلوغ.\n\n## اشیاء\n\n- ماشین\n- درخت",
                {"prompt_tokens": 100, "completion_tokens": 20,
                 "total_tokens": 120})

    monkeypatch.setattr(
        "app.ai.openai_compat.OpenAICompatProvider.describe_image_with_usage",
        fake_describe)
    doc = F.parse_image_meta(_png(), "street.png")
    assert doc.parse_meta["parser"] == "image-structured"
    assert doc.parse_meta["usage"]["total_tokens"] == 120
    assert doc.parse_meta["model"]
    kinds = [e.kind for e in doc.elements]
    assert kinds[0] == "heading" and "list" in kinds
    assert doc.elements[1].metadata["section"] == "صحنه"


def test_image_failure_is_metadata_only():
    from app.knowledge.parsers.fallback import parse_image_meta

    def boom(blob):
        raise RuntimeError("vision down")

    doc = parse_image_meta(_png(), "x.png", describe_fn=boom)
    assert doc.pages[0].text.startswith("[image 320x200")
    assert doc.elements == []
    bad = parse_image_meta(b"not-an-image", "x.bin", describe_fn=boom)
    assert bad.pages and bad.parse_meta["parser"] == "image-meta"


def test_image_ingest_sectioned_chunks(db, monkeypatch):
    import workers.tasks as T
    from app.core.config import get_settings

    monkeypatch.setattr(get_settings(), "DOCUMENT_VISION_ENABLED", False)
    _, w, kb = _seed_ws(db)
    st = FakeStorage()
    src = Source(id=uuid.uuid4(), workspace_id=w, kb_id=kb.id, type="image",
                 filename="pic.png", storage_key="k", status="pending")
    db.add(src)
    db.commit()
    st.put("k", _png())
    res = run_ingest(src.id, workspace_id=w, db=db, storage=st,
                     embed_fn=fake_embed,
                     describe_fn=lambda blob: "## صحنه\n\nخیابان.")
    assert res["ok"] is True
    chunks = db.query(Chunk).all()
    assert chunks and any("خیابان" in c.content for c in chunks)


# ---------- audio chapters ----------

def test_chapterize_gaps_and_titles():
    from app.knowledge.parsers.audio import chapterize

    segs = [(0, 30_000, "سلام به کلاس امروز درباره فیزیک صحبت می‌کنیم"),
            (30_000, 60_000, "ادامه بحث فصل اول"),
            (300_000, 330_000, "حالا می‌رویم سراغ فصل دوم شیمی"),
            (330_000, 360_000, "ادامه فصل دوم")]
    chs = chapterize(segs)
    assert len(chs) == 2
    assert chs[0]["start_ms"] == 0 and chs[0]["end_ms"] == 60_000
    assert chs[1]["start_ms"] == 300_000 and chs[1]["end_ms"] == 360_000
    assert chs[1]["title"].startswith("حالا می‌رویم")
    assert all(set(c) == {"start_ms", "end_ms", "title"} for c in chs)
    # no speaker labels anywhere: diarization is out of scope
    assert all("speaker" not in str(c).lower() for c in chs)


def test_chapterize_merges_short_and_handles_untimed():
    from app.knowledge.parsers.audio import chapterize

    segs = [(0, 200_000, "بخش طولانی اول که بیش از دو دقیقه طول می‌کشد و ادامه دارد"),
            (500_000, 510_000, "میان‌پرده کوتاه"),
            (900_000, 1_100_000, "بخش طولانی دوم با محتوای کافی و ادامه‌دار")]
    chs = chapterize(segs)
    assert len(chs) == 2  # the 10s interstitial merges, never dropped
    assert chs[0]["end_ms"] == 510_000
    assert chapterize([(0, 0, "متن بدون زمان‌بندی")])[0]["end_ms"] == 0
    assert chapterize([]) == []
    assert chapterize([(0, 0, "   ")]) == []


def test_dispatcher_audio_attaches_chapters():
    from app.knowledge import parsers as P

    def fake_transcribe(blob, filename):
        return [(0, 30_000, "مقدمه جلسه امروز"),
                (400_000, 430_000, "بخش دوم جلسه درباره بودجه")]

    doc = P.parse_source("audio", "m.mp3", b"ID3",
                         transcribe_fn=fake_transcribe)
    assert len(doc.pages) == 2
    kinds = [e.kind for e in doc.elements]
    assert kinds == ["chapter", "chapter"]
    assert doc.elements[1].metadata["start_ms"] == 400_000
    assert "بودجه" in doc.elements[1].metadata["title"]


def test_audio_windows_carry_section():
    from app.knowledge.chunking import chunk_parsed
    from app.knowledge.parsers.base import ParsedDocument, ParsedPage

    pages = [ParsedPage(start_ms=i * 30_000, end_ms=(i + 1) * 30_000,
                        text=f"متن بخش شماره {i} درباره موضوع جلسه")
             for i in range(8)]  # 4 minutes, one gap-free chapter
    doc = ParsedDocument(title="m", pages=pages)
    from app.knowledge.parsers.audio import chapterize
    from app.knowledge.parsers.base import ParsedElement

    chs = chapterize([(p.start_ms, p.end_ms, p.text) for p in pages])
    assert len(chs) == 1
    doc.elements = [ParsedElement(
        element_id="ch1", kind="chapter", reading_order=1,
        text_markdown=chs[0]["title"],
        metadata={"kind": "chapter", "element_id": "ch1",
                  "start_ms": chs[0]["start_ms"],
                  "end_ms": chs[0]["end_ms"], "title": chs[0]["title"]})]
    drafts = chunk_parsed(doc, is_audio=True)
    assert drafts
    assert all(d.metadata.get("section") for d in drafts)
    assert all(d.metadata.get("chapter_start_ms") == 0 for d in drafts)
    assert all(d.start_ms is not None and d.end_ms is not None for d in drafts)


def test_transcript_includes_chapters(db):
    from app.knowledge.models import Document
    from app.knowledge.service import source_transcript

    _, w, kb = _seed_ws(db)
    src = Source(id=uuid.uuid4(), workspace_id=w, kb_id=kb.id, type="audio",
                 filename="m.mp3", status="ready")
    db.add(src)
    db.flush()
    doc = Document(id=uuid.uuid4(), workspace_id=w, source_id=src.id,
                   title="m")
    db.add(doc)
    db.flush()
    db.add(PageSegment(id=uuid.uuid4(), workspace_id=w, document_id=doc.id,
                       start_ms=0, end_ms=30_000, text="مقدمه جلسه"))
    db.add(PageSegment(id=uuid.uuid4(), workspace_id=w, document_id=doc.id,
                       start_ms=500_000, end_ms=530_000,
                       text="بخش دوم درباره بودجه"))
    db.commit()
    out = source_transcript(db, src)
    assert len(out["segments"]) == 2
    assert len(out["chapters"]) == 2
    assert out["chapters"][1]["start_ms"] == 500_000
    assert "بودجه" in out["chapters"][1]["title"]


def test_audio_ingest_chunks_carry_section(db):
    _, w, kb = _seed_ws(db)
    st = FakeStorage()
    src = Source(id=uuid.uuid4(), workspace_id=w, kb_id=kb.id, type="audio",
                 filename="m.mp3", storage_key="k", status="pending")
    db.add(src)
    db.commit()
    st.put("k", b"ID3")

    def fake_transcribe(blob, filename):
        return [(0, 30_000, "مقدمه جلسه امروز درباره پروژه"),
                (30_000, 60_000, "ادامه بحث پروژه و وظایف"),
                (600_000, 630_000, "بخش پایانی و جمع‌بندی تصمیم‌ها")]

    res = run_ingest(src.id, workspace_id=w, db=db, storage=st,
                     embed_fn=fake_embed, transcribe_fn=fake_transcribe)
    assert res["ok"] is True
    chunks = db.query(Chunk).all()
    assert chunks
    assert any(c.chunk_metadata.get("section") for c in chunks)
    assert any(c.chunk_metadata.get("chapter_start_ms", 0) >= 600_000
               for c in chunks)


# ---------- excel relations ----------

def _rel_workbook() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["order_id", "customer_id", "amount"])
    ws.append([1, "C1", 100])
    ws.append([2, "C2", 200])
    ws.append([3, "C1", 50])
    ws2 = wb.create_sheet("Customers")
    ws2.append(["customer_id", "name"])
    ws2.append(["C1", "Ali"])
    ws2.append(["C2", "Sara"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_detect_relations_shared_column():
    from app.knowledge.excel.parse import detect_relations, parse_workbook

    sheets = parse_workbook(_rel_workbook())
    rels = detect_relations(sheets)
    assert len(rels) == 1
    r = rels[0]
    assert r["kind"] == "shared-column"
    assert {r["from_table"], r["to_table"]} == {"orders", "customers"}
    assert r["from_col"] == "customer_id" and r["to_col"] == "customer_id"
    assert r["shared_values"] >= 2


def test_detect_relations_negatives():
    from openpyxl import Workbook

    from app.knowledge.excel.parse import detect_relations, parse_workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "A"
    ws.append(["id", "item"])
    ws.append([1, "x"])
    ws2 = wb.create_sheet("B")
    ws2.append(["id", "cost"])
    ws2.append([9, 5])  # same col name, disjoint values -> no relation
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    sheets = parse_workbook(buf.getvalue())
    assert detect_relations(sheets) == []
    assert detect_relations(parse_workbook(_single_sheet())) == []


def _single_sheet() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Only"
    ws.append(["a"])
    ws.append([1])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_detect_relations_value_overlap():
    from app.knowledge.excel.parse import SheetInfo, detect_relations

    def sh(name, cols, rows):
        from app.knowledge.excel.parse import ExcelColumn

        return SheetInfo(
            name=name, table=name.lower(),
            columns=[ExcelColumn(name=c, col=c, dtype="text") for c in cols],
            n_rows=len(rows), sample_rows=[], rows=rows)

    sheets = [sh("S1", ["code_a"], [[f"v{i}"] for i in range(10)]),
              sh("S2", ["code_b"], [[f"v{i}"] for i in range(10)])]
    rels = detect_relations(sheets)
    assert len(rels) == 1 and rels[0]["kind"] == "value-overlap"


def test_relations_in_summary_and_parse_meta(db):
    from app.knowledge.parsers.excel import parse_excel

    doc = parse_excel(_rel_workbook())
    assert doc.parse_meta["relations"]
    assert any("Relationships" in p.text for p in doc.pages)


def test_validate_join_accept_reject():
    from app.knowledge.excel.store import validate_sql

    allow = {"orders": {"order_id", "customer_id", "amount"},
             "customers": {"customer_id", "name"}}
    rels = [{"from_table": "orders", "from_col": "customer_id",
             "to_table": "customers", "to_col": "customer_id",
             "kind": "shared-column"}]
    ok = validate_sql(
        "SELECT customers.name, SUM(orders.amount) FROM orders "
        "JOIN customers ON orders.customer_id = customers.customer_id "
        "GROUP BY customers.name", allow, rels)
    assert ok["join"] == "customers"
    assert "JOIN" in ok["sql"] and "GROUP BY" in ok["sql"]
    # unknown edge rejected
    with pytest.raises(ValueError):
        validate_sql("SELECT a FROM orders JOIN customers "
                     "ON orders.amount = customers.name", allow, rels)
    # ambiguous bare column rejected
    with pytest.raises(ValueError):
        validate_sql("SELECT customer_id FROM orders "
                     "JOIN customers ON orders.customer_id = customers.customer_id",
                     allow, rels)
    # joins without relations rejected
    with pytest.raises(ValueError):
        validate_sql("SELECT a FROM orders JOIN customers "
                     "ON orders.customer_id = customers.customer_id", allow)
    # multi-join rejected (2-table INNER JOIN only)
    with pytest.raises(ValueError):
        validate_sql("SELECT a FROM orders JOIN customers "
                     "ON orders.customer_id = customers.customer_id "
                     "JOIN x ON x.a = orders.amount", allow, rels)
    # legacy single-table shape byte-identical
    legacy = validate_sql(
        "SELECT region, SUM(amount) FROM sales GROUP BY region LIMIT 5",
        {"sales": {"region", "amount"}})
    assert legacy["sql"] == (
        'SELECT region, SUM(amount) FROM "sales" '
        'GROUP BY "region" LIMIT 5')
    assert legacy["join"] is None


def _rel_source(db, storage, wid, kb_id) -> Source:
    src = Source(id=uuid.uuid4(), workspace_id=wid, kb_id=kb_id, type="excel",
                 filename="shop.xlsx", status="pending")
    db.add(src)
    db.flush()
    key = f"k/{src.id}/shop.xlsx"
    storage.put(key, _rel_workbook())
    src.storage_key = key
    db.commit()
    return src


def test_excel_join_end_to_end(db):
    from app.agents.tools import ToolContext, excel_query

    _, w, kb = _seed_ws(db)
    st = FakeStorage()
    src = _rel_source(db, st, w, kb.id)
    res = run_ingest(src.id, db=db, storage=st, embed_fn=fake_embed)
    assert res["ok"] is True
    db.refresh(src)
    rels = (src.parse_meta or {}).get("relations") or []
    assert len(rels) == 1
    ctx = ToolContext(db=db, workspace_id=w, user_id=uuid.uuid4(),
                      storage=st, kb_ids=[kb.id])
    out = excel_query(ctx, "مجموع amount به تفکیک name")
    assert len(out) == 1
    got = {r[0]: r[1] for r in out[0]["rows"]}
    assert got == {"Ali": "150", "Sara": "200"}
    assert "JOIN" in out[0]["sql"]

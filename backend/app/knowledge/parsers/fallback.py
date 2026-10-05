"""Local parsers: deterministic, dependency-light, always available.

Covers pdf/docx/pptx/txt/md/url/note. pdf is the flat safety net under
Gemini structured extraction; docx keeps tables in document order;
txt/md/note gain heading/paragraph/list structure. Selected in
parsers/__init__.py.
"""

import io
import json

from app.knowledge.parsers.base import ParsedDocument, ParsedPage


def parse_pdf(blob: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(blob))
    pages = []
    for i, page in enumerate(reader.pages):
        pages.append(ParsedPage(page_no=i + 1, text=(page.extract_text() or "").strip()))
    return ParsedDocument(pages=pages, page_count=len(pages))


def parse_docx(blob: bytes) -> ParsedDocument:
    """Word paragraphs AND tables in document order.

    Tables were previously dropped entirely (in shenakht.docx ~26% of the
    words live inside a table), so large Word files were learned
    incompletely. Cells of a row are joined with " | ", rows with newlines.
    """
    from docx import Document as DocxDocument
    from docx.table import Table as _DocxTable
    from docx.text.paragraph import Paragraph as _DocxParagraph

    doc = DocxDocument(io.BytesIO(blob))
    parts: list[str] = []
    try:
        body = doc.element.body
    except Exception:
        body = None
    if body is not None:
        for child in body.iterchildren():
            tag = str(getattr(child, "tag", ""))
            try:
                if tag.endswith("}p"):
                    t = _DocxParagraph(child, doc).text.strip()
                    if t:
                        parts.append(t)
                elif tag.endswith("}tbl"):
                    rows = []
                    for row in _DocxTable(child, doc).rows:
                        cells = [(c.text or "").strip() for c in row.cells]
                        line = " | ".join(cells).strip(" |").strip()
                        if line:
                            rows.append(line)
                    if rows:
                        parts.append("\n".join(rows))
            except Exception:
                continue
    if not parts:  # odd document structure: legacy paragraph-only path
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        if text.strip():
            parts.append(text.strip())
    return ParsedDocument(pages=[ParsedPage(text="\n".join(parts).strip())], page_count=0)


def parse_pptx(blob: bytes) -> ParsedDocument:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(blob))
    pages = []
    for i, slide in enumerate(prs.slides):
        parts = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
        pages.append(ParsedPage(page_no=i + 1, text="\n".join(parts).strip()))
    return ParsedDocument(pages=pages, page_count=len(pages))


def parse_text(blob: bytes) -> ParsedDocument:
    return ParsedDocument(pages=[ParsedPage(text=blob.decode("utf-8", errors="replace").strip())])


def parse_structured_text(blob: bytes, title: str = "") -> ParsedDocument:
    """txt/md/note: same single-page legacy view as parse_text, PLUS
    structured elements (heading/paragraph/list + section chain) so
    chunking/RAG sees document structure, not just flat text."""
    from app.knowledge.parsers.normalize import structure_plain_text

    text = blob.decode("utf-8", errors="replace").strip()
    doc = structure_plain_text(text, title=title)
    doc.pages = [ParsedPage(text=text)]
    return doc


def parse_url_descriptor(blob: bytes, fetch_fn=None,
                         transcribe_fn=None,
                         provider_overrides: dict | None = None) -> ParsedDocument:
    """Provider-routed URL extraction (YouTube/Instagram/Aparat/generic).

    Video URLs yield timestamped transcript pages (same shape as audio);
    text URLs yield a single text page. Descriptor fallback on any failure
    so the job never fails just because a site is down.
    """
    data = json.loads(blob.decode("utf-8"))
    title = str(data.get("title") or data.get("url") or "url")
    url = str(data.get("url") or "")
    provider_err = ""
    provider = "generic"
    if url:
        try:
            from app.knowledge.parsers.providers import detect, extract

            try:
                provider = detect(url)
            except Exception:
                provider = "generic"
            got = extract(url, fetch_fn=fetch_fn,
                          transcribe_fn=transcribe_fn,
                          provider_overrides=provider_overrides)
            meta = {"parser": "url-provider", "provider": got.provider,
                    "url": got.url}
            if got.title:
                meta["provider_title"] = got.title[:300]
            if got.kind == "transcript" and got.segments:
                duration = max((e for _, e, _ in got.segments), default=0)
                return ParsedDocument(
                    title=got.title or title,
                    pages=[ParsedPage(start_ms=s, end_ms=e, text=t)
                           for s, e, t in got.segments],
                    duration_ms=duration, parse_meta=meta)
            if (got.text or "").strip():
                meta.update(got.meta or {})
                return ParsedDocument(title=got.title or title,
                                      pages=[ParsedPage(text=got.text.strip())],
                                      parse_meta=meta)
        except Exception as exc:
            # Transparent fallback: the processing view shows WHICH provider
            # failed and why, instead of a silent descriptor-only source.
            provider_err = f"{provider}: {type(exc).__name__}: {exc}"[:300]
    text = f"{title}\n{url}".strip()
    meta = {"parser": "url-descriptor-fallback"}
    if provider_err:
        meta["provider_error"] = provider_err
    return ParsedDocument(title=title, pages=[ParsedPage(text=text)],
                          parse_meta=meta)


def parse_note_descriptor(blob: bytes) -> ParsedDocument:
    data = json.loads(blob.decode("utf-8"))
    title = str(data.get("title") or "note")
    content = str(data.get("content", "") or "").strip()
    from app.knowledge.parsers.normalize import structure_plain_text

    doc = structure_plain_text(content, title=title)
    doc.title = title
    doc.pages = [ParsedPage(text=content)]
    return doc


IMAGE_STRUCT_PROMPT = (
    "این تصویر را به فارسی، ساختاریافته توصیف کن. فقط همین بخش‌ها را برگردان "
    "(بخشی که در تصویر نیست را حذف کن، چیزی حدس نزن):\n"
    "## صحنه\n(یک پاراگراف: چه چیزی، کجا، چه می‌کند)\n"
    "## اشیاء\n(فهرست مهم‌ترین اشیاء قابل مشاهده)\n"
    "## متن داخل تصویر\n(متن‌های خوانا دقیقاً همان‌طور که نوشته شده‌اند؛ اگر متنی نیست بنویس «بدون متن»)\n"
    "## رنگ‌ها و ترکیب‌بندی\n(یک پاراگراف کوتاه)"
)


def _default_image_model() -> str | None:
    """Vision model for image structuring: admin `vision.provider` default
    (no-DB merge, so panel customizations need the worker vision stage
    which resolves per-source with a session)."""
    try:
        from app.ai.settings import get_setting

        model = (get_setting(None, "vision.provider") or {}).get("model")
        if (model or "").strip():
            return str(model).strip()
    except Exception:
        pass
    return None


def parse_image_meta(blob: bytes, filename: str, describe_fn=None,
                     model: str | None = None) -> ParsedDocument:
    """Structured image extraction: dimensions via pillow + sectioned
    visual description (scene/objects/visible-text/colors) as heading/
    paragraph/list elements, so image chunks carry sections like every
    other source. `describe_fn` (str-returning seam, mocked in tests)
    keeps its contract; the default path uses the usage-aware vision
    call and records model+tokens in parse_meta (feeds the P3 ledger
    through the worker's extraction event). Never fails the job:
    worst case the segment carries metadata only."""
    from app.knowledge.parsers.normalize import structure_plain_text

    meta = ""
    try:
        from PIL import Image

        with Image.open(io.BytesIO(blob)) as im:
            meta = f"[image {im.width}x{im.height} {im.format or ''} {filename}]".strip()
    except Exception:
        meta = f"[image {filename}]"
    description, usage, used_model = "", {}, None
    if describe_fn is not None:
        try:
            description = (describe_fn(blob) or "").strip()
        except Exception:
            description = ""
    else:
        try:
            from app.ai.openai_compat import OpenAICompatProvider

            used_model = model or _default_image_model()
            description, usage = OpenAICompatProvider().describe_image_with_usage(
                blob, IMAGE_STRUCT_PROMPT, model=used_model, max_tokens=800)
            description = (description or "").strip()
        except Exception:
            description = ""
    doc = structure_plain_text(description, title=filename) \
        if description else ParsedDocument(title=filename)
    doc.pages = [ParsedPage(text="\n".join(p for p in (meta, description) if p))]
    doc.parse_meta = {"parser": "image-structured" if usage else "image-meta",
                      "image_meta": meta}
    if used_model:
        doc.parse_meta["model"] = used_model
    if usage:
        doc.parse_meta["usage"] = {
            "model": used_model or "default", "calls": 1,
            "prompt_tokens": usage.get("prompt_tokens", 0),
            "completion_tokens": usage.get("completion_tokens", 0),
            "total_tokens": usage.get("total_tokens", 0)}
    return doc


FALLBACK = {
    "pdf": parse_pdf,
    "docx": parse_docx,
    "pptx": parse_pptx,
    "txt": parse_structured_text,
    "md": parse_structured_text,
    "url": parse_url_descriptor,
    "note": parse_note_descriptor,
}
